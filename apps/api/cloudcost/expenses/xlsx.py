"""Lectura segura de planillas Excel (.xlsx) como texto CSV para los analizadores de gastos.

Solo se leen VALORES: las fórmulas se usan por su último resultado calculado (nunca se evalúan) y el texto de una celda jamás se interpreta como
código. Se rechazan los libros con macros (.xlsm / vbaProject), los que no son un .xlsx válido y los que se expanden de forma sospechosa
(bomba de compresión). Cada hoja con datos se devuelve como un «archivo» aparte.
"""
from __future__ import annotations

import csv
import io
import zipfile
from datetime import date, datetime, time
from decimal import Decimal

from .parser import MAX_ROWS, ExpenseFormatError

MAX_XLSX_BYTES = 9_000_000                  # tamaño del archivo comprimido
MAX_UNCOMPRESSED_BYTES = 60_000_000         # suma de lo descomprimido: más que eso es una bomba o un libro fuera de alcance
MAX_ZIP_ENTRIES = 2_000
MAX_SHEETS = 12
MAX_COLUMNS = 60
ZIP_MAGIC = b"PK\x03\x04"


def _cell_text(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0, 0) else v.isoformat(timespec="seconds")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.isoformat(timespec="seconds")
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v != v or v in (float("inf"), float("-inf")):
            return ""
        d = Decimal(repr(v))
        return format(d.quantize(Decimal(1)) if d == d.to_integral_value() else d.normalize(), "f")
    return str(v).replace("\r", " ").replace("\n", " ").strip()


def _check_container(data: bytes) -> None:
    if len(data) > MAX_XLSX_BYTES:
        raise ExpenseFormatError(f"El Excel supera el máximo de {MAX_XLSX_BYTES // 1_000_000} MB")
    if not data.startswith(ZIP_MAGIC):
        raise ExpenseFormatError("No es un archivo .xlsx válido (los .xls antiguos no se admiten: guárdalo como .xlsx o CSV)")
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        infos = zf.infolist()
    except zipfile.BadZipFile as exc:
        raise ExpenseFormatError("El archivo .xlsx está dañado") from exc
    if len(infos) > MAX_ZIP_ENTRIES:
        raise ExpenseFormatError("El .xlsx tiene demasiados componentes")
    names = {i.filename.lower() for i in infos}
    if any(n.endswith("vbaproject.bin") or n.startswith("xl/macrosheets/") for n in names):
        raise ExpenseFormatError("El libro contiene macros: guárdalo como .xlsx sin macros o como CSV")
    if "xl/workbook.xml" not in names:
        raise ExpenseFormatError("No es un libro de Excel (.xlsx) válido")
    if sum(i.file_size for i in infos) > MAX_UNCOMPRESSED_BYTES:
        raise ExpenseFormatError("El .xlsx se expande a un tamaño excesivo y no se procesa")


def xlsx_to_csv_texts(data: bytes, filename: str = "") -> list[tuple[str, str]]:
    """[(nombre, texto CSV con «;»)] — una entrada por hoja con datos. El nombre lleva la hoja cuando hay varias."""
    _check_container(data)
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except ImportError as exc:                                         # pragma: no cover — openpyxl está en requirements.txt
        raise ExpenseFormatError("Este servidor no puede leer Excel; sube un CSV") from exc
    except Exception as exc:                                           # noqa: BLE001 — openpyxl lanza tipos muy variados ante archivos dañados
        raise ExpenseFormatError("No se pudo abrir el .xlsx (¿está protegido con contraseña o dañado?)") from exc
    out: list[tuple[str, str]] = []
    try:
        sheets = [ws for ws in wb.worksheets if getattr(ws, "sheet_state", "visible") == "visible"][:MAX_SHEETS]
        for ws in sheets:
            buf = io.StringIO()
            writer = csv.writer(buf, delimiter=";", lineterminator="\n")
            count = 0
            for row in ws.iter_rows(values_only=True):
                cells = [_cell_text(c) for c in row[:MAX_COLUMNS]]
                while cells and not cells[-1]:
                    cells.pop()
                writer.writerow(cells)
                count += 1
                if count > MAX_ROWS + 200:
                    raise ExpenseFormatError(f"La hoja «{ws.title}» supera el máximo de {MAX_ROWS} filas")
            text = buf.getvalue()
            if text.replace(";", "").strip():
                out.append((ws.title, text))
    finally:
        wb.close()
    if not out:
        raise ExpenseFormatError("El Excel no tiene hojas con datos")
    base = filename or "libro.xlsx"
    return [(base if len(out) == 1 else f"{base} · {title}", text) for title, text in out]
