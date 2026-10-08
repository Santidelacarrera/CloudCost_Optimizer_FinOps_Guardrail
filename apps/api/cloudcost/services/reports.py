"""Reportes ejecutivos de desperdicio y ahorro proyectado (PDF y Excel) para la gerencia.

Separación deliberada: ``load_report_data`` lee de la base (dentro de la transacción del tenant) y los renderizadores
son funciones puras sobre ``ReportData``, así se prueban sin base de datos y el mismo contenido sale en ambos formatos.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Any

from psycopg import Connection

OPEN_STATUSES = ("PENDING_APPROVAL", "APPROVED", "PR_CREATED", "MERGED", "DEPLOYED")
RULE_LABELS = {
    "ec2_downsize": "Instancias sobredimensionadas",
    "ec2_idle": "Instancias ociosas",
    "ebs_orphan": "Volúmenes huérfanos",
    "snapshot_old": "Snapshots antiguos",
    "rds_idle": "Bases de datos sin uso",
}
ENV_LABELS = {"production": "Producción", "staging": "Staging", "development": "Desarrollo", "test": "Pruebas", "unknown": "Sin clasificar"}
STATUS_LABELS = {"PENDING_APPROVAL": "Pendiente de aprobación", "APPROVED": "Aprobada", "PR_CREATED": "PR creado",
                 "MERGED": "PR fusionado", "DEPLOYED": "Desplegada", "VERIFIED": "Verificada"}
REAL_SOURCES = {"cost_explorer", "cost_explorer_tag", "cur", "import"}


@dataclass
class ReportData:
    organization: str
    generated_at: datetime
    monthly_spend: float
    opportunities: list[dict[str, Any]] = field(default_factory=list)   # ordenadas por ahorro descendente
    daily_cost: list[tuple[str, float]] = field(default_factory=list)   # (YYYY-MM-DD, USD)
    verified: list[dict[str, Any]] = field(default_factory=list)

    # ---- agregados (se derivan de las oportunidades, así Excel y PDF nunca discrepan)
    @property
    def potential_monthly(self) -> float:
        return round(sum(o["monthly_savings"] for o in self.opportunities), 2)

    @property
    def potential_annual(self) -> float:
        return round(self.potential_monthly * 12, 2)

    @property
    def savings_pct(self) -> float:
        return round(self.potential_monthly / self.monthly_spend * 100, 1) if self.monthly_spend > 0 else 0.0

    @property
    def verified_monthly(self) -> float:
        return round(sum(v["observed"] for v in self.verified), 2)

    def grouped(self, key: str) -> list[tuple[str, int, float]]:
        acc: dict[str, list[float]] = {}
        for o in self.opportunities:
            acc.setdefault(o[key], []).append(o["monthly_savings"])
        return sorted(((k, len(v), round(sum(v), 2)) for k, v in acc.items()), key=lambda t: -t[2])

    @property
    def real_cost_share(self) -> tuple[int, int]:
        real = sum(1 for o in self.opportunities if o["cost_source"] in REAL_SOURCES)
        return real, len(self.opportunities)


# --------------------------------------------------------------------------- lectura
def load_report_data(conn: Connection, *, now: datetime | None = None) -> ReportData:
    org = conn.execute("select name from organizations limit 1").fetchone()
    spend = conn.execute("select coalesce(sum(monthly_cost), 0) as v from resources where active").fetchone()["v"]
    rows = conn.execute(
        """select r.title, r.rule_id, r.action, r.risk, r.priority, r.status, r.confidence, r.destructive,
                  r.estimated_monthly_savings as savings, res.resource_id, res.environment,
                  coalesce(r.evidence -> 'cost_basis' ->> 'source', res.cost_source) as cost_source
             from recommendations r join resources res on res.id = r.resource_pk
            where r.status = any(%s) order by r.estimated_monthly_savings desc, r.title""", (list(OPEN_STATUSES),)).fetchall()
    daily = conn.execute(
        """select usage_date, sum(amount) as amount from cost_records
            where usage_date >= current_date - 30 group by usage_date order by usage_date""").fetchall()
    verified = conn.execute(
        """select r.title, v.expected_monthly_savings as expected, v.observed_monthly_savings as observed,
                  v.realization_pct as pct from savings_verifications v join recommendations r on r.id = v.recommendation_id
            order by v.created_at desc limit 50""").fetchall()
    return ReportData(
        organization=(org or {}).get("name") or "Organización",
        generated_at=now or datetime.now(timezone.utc),
        monthly_spend=float(spend),
        opportunities=[{
            "title": r["title"], "resource_id": r["resource_id"], "rule": RULE_LABELS.get(r["rule_id"], r["rule_id"]),
            "action": r["action"], "environment": ENV_LABELS.get(r["environment"], r["environment"]), "risk": r["risk"],
            "priority": r["priority"], "monthly_savings": float(r["savings"]), "confidence": float(r["confidence"]),
            "status": STATUS_LABELS.get(r["status"], r["status"]), "cost_source": r["cost_source"] or "estimate",
            "destructive": bool(r["destructive"])} for r in rows],
        daily_cost=[(str(d["usage_date"]), round(float(d["amount"]), 2)) for d in daily],
        verified=[{"title": v["title"], "expected": float(v["expected"]), "observed": float(v["observed"]),
                   "pct": float(v["pct"]) if v["pct"] is not None else None} for v in verified])


_MONTHS = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def long_date(d: datetime) -> str:
    return f"{d.day} de {_MONTHS[d.month - 1]} de {d.year}"


def usd(x: float) -> str:
    return f"USD {x:,.2f}"


def source_label(source: str) -> str:
    return {"cost_explorer": "Facturado (Cost Explorer)", "cost_explorer_tag": "Facturado, repartido por etiqueta",
            "cur": "Facturado (CUR)", "import": "Importado"}.get(source, "Estimado")


def headline(data: ReportData) -> str:
    if not data.opportunities:
        return "No hay oportunidades de ahorro abiertas en este momento."
    top = data.opportunities[0]
    return (f"Hay {len(data.opportunities)} oportunidades abiertas por {usd(data.potential_monthly)} al mes "
            f"({usd(data.potential_annual)} al año, {data.savings_pct:g}% del gasto). La mayor es «{top['title']}» "
            f"con {usd(top['monthly_savings'])} al mes.")


METHOD_NOTES = [
    "Cada oportunidad es una propuesta: nada se cambia en la nube. Los cambios llegan como Pull Request y requieren aprobación humana.",
    "«Estimado» significa que el costo sale de una tabla de precios; «Facturado» viene de AWS Cost Explorer y es el dato recomendado "
    "para decidir.",
    "Los snapshots se valoran como cota superior (son incrementales); el ahorro real puede ser menor.",
    "El ahorro verificado es el medido después del despliegue (costo base menos costo observado), no una proyección.",
]


# --------------------------------------------------------------------------- Excel
def render_xlsx(data: ReportData) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ink, accent, soft = "1F2A37", "2F6F5E", "EEF4F1"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=accent)
    thin = Side(style="thin", color="D5DDD9")
    money = '"USD" #,##0.00'

    def header(ws, row: int, labels: list[str]) -> None:
        for c, label in enumerate(labels, start=1):
            cell = ws.cell(row=row, column=c, value=label)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    def widths(ws, values: list[int]) -> None:
        for i, w in enumerate(values, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()
    summary, opp = wb.active, wb.create_sheet("Oportunidades")
    summary.title = "Resumen"

    # ---- Oportunidades (la fuente de todas las fórmulas)
    header(opp, 1, ["Hallazgo", "Recurso", "Categoría", "Acción", "Entorno", "Riesgo", "Prioridad", "Ahorro mensual (USD)",
                    "Ahorro anual (USD)", "Confianza", "Estado", "Origen del costo"])
    for i, o in enumerate(data.opportunities, start=2):
        opp.append([o["title"], o["resource_id"], o["rule"], o["action"], o["environment"], o["risk"], o["priority"],
                    o["monthly_savings"], f"=H{i}*12", o["confidence"], o["status"], source_label(o["cost_source"])])
        opp[f"H{i}"].number_format = opp[f"I{i}"].number_format = money
        opp[f"J{i}"].number_format = "0%"
    last = max(2, len(data.opportunities) + 1)
    opp.freeze_panes = "A2"
    opp.auto_filter.ref = f"A1:L{last}"
    widths(opp, [52, 26, 30, 20, 16, 10, 10, 20, 18, 11, 24, 32])

    # ---- Resumen
    summary["A1"], summary["A1"].font = "Reporte de desperdicio y ahorro proyectado", Font(bold=True, size=16, color=ink)
    summary["A2"] = f"{data.organization} · generado el {data.generated_at:%d-%m-%Y %H:%M} UTC"
    summary["A2"].font = Font(color="5B6770")
    summary["A4"] = headline(data)
    summary["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    summary.merge_cells("A4:D5")
    kpis = [
        ("Gasto mensual actual", data.monthly_spend, money),
        ("Ahorro mensual potencial", f"=SUM(Oportunidades!H2:H{last})", money),
        ("Ahorro anual proyectado", "=B8*12", money),
        ("Ahorro sobre el gasto", "=IF(B7>0,B8/B7,0)", "0.0%"),
        ("Oportunidades abiertas", f"=COUNTA(Oportunidades!A2:A{last})", "0"),
        ("Ahorro mensual verificado (medido)", "=SUM('Ahorro verificado'!C2:C500)", money),
    ]
    for r, (label, value, fmt) in enumerate(kpis, start=7):
        summary.cell(row=r, column=1, value=label).font = Font(bold=True, color=ink)
        c = summary.cell(row=r, column=2, value=value)
        c.number_format, c.font = fmt, Font(bold=True, size=12, color=accent)
        for col in (1, 2):
            summary.cell(row=r, column=col).fill = PatternFill("solid", fgColor=soft)
            summary.cell(row=r, column=col).border = Border(bottom=thin)
    widths(summary, [38, 24, 18, 18])

    # ---- Por categoría / entorno (SUMIF sobre Oportunidades: si editas una fila, todo se recalcula)
    cat = wb.create_sheet("Por categoría")
    header(cat, 1, ["Categoría", "Oportunidades", "Ahorro mensual (USD)", "Ahorro anual (USD)"])
    cats = data.grouped("rule")
    for i, (name, _n, _s) in enumerate(cats, start=2):
        cat.append([name, f"=COUNTIF(Oportunidades!C2:C{last},A{i})", f"=SUMIF(Oportunidades!C2:C{last},A{i},Oportunidades!H2:H{last})", f"=C{i}*12"])
        cat[f"C{i}"].number_format = cat[f"D{i}"].number_format = money
    base = len(cats) + 4
    header(cat, base, ["Entorno", "Oportunidades", "Ahorro mensual (USD)", "Ahorro anual (USD)"])
    envs = data.grouped("environment")
    for i, (name, _n, _s) in enumerate(envs, start=base + 1):
        cat.append([name, f"=COUNTIF(Oportunidades!E2:E{last},A{i})", f"=SUMIF(Oportunidades!E2:E{last},A{i},Oportunidades!H2:H{last})", f"=C{i}*12"])
        cat[f"C{i}"].number_format = cat[f"D{i}"].number_format = money
    widths(cat, [34, 16, 22, 20])
    if cats:
        chart = BarChart()
        chart.title, chart.y_axis.title, chart.type = "Ahorro mensual por categoría", "USD", "bar"
        chart.add_data(Reference(cat, min_col=3, min_row=1, max_row=len(cats) + 1), titles_from_data=True)
        chart.set_categories(Reference(cat, min_col=1, min_row=2, max_row=len(cats) + 1))
        chart.legend, chart.height, chart.width = None, 8, 18
        cat.add_chart(chart, "F2")

    # ---- Gasto diario
    daily = wb.create_sheet("Gasto diario")
    header(daily, 1, ["Fecha", "Gasto (USD)"])
    for d, amount in data.daily_cost:
        daily.append([d, amount])
    for row in range(2, len(data.daily_cost) + 2):
        daily[f"B{row}"].number_format = money
    widths(daily, [14, 18])
    if len(data.daily_cost) > 1:
        line = LineChart()
        line.title, line.legend, line.height, line.width = "Gasto diario (últimos 30 días)", None, 8, 18
        line.add_data(Reference(daily, min_col=2, min_row=1, max_row=len(data.daily_cost) + 1), titles_from_data=True)
        line.set_categories(Reference(daily, min_col=1, min_row=2, max_row=len(data.daily_cost) + 1))
        daily.add_chart(line, "D2")

    # ---- Ahorro verificado
    ver = wb.create_sheet("Ahorro verificado")
    header(ver, 1, ["Recomendación", "Ahorro esperado (USD)", "Ahorro observado (USD)", "Realización"])
    for i, v in enumerate(data.verified, start=2):
        ver.append([v["title"], v["expected"], v["observed"], f"=IF(B{i}>0,C{i}/B{i},0)"])
        ver[f"B{i}"].number_format = ver[f"C{i}"].number_format = money
        ver[f"D{i}"].number_format = "0%"
    widths(ver, [60, 22, 22, 14])

    notes = wb.create_sheet("Notas")
    notes["A1"], notes["A1"].font = "Cómo leer este reporte", Font(bold=True, size=13, color=ink)
    for i, text in enumerate(METHOD_NOTES, start=3):
        notes.cell(row=i, column=1, value=text).alignment = Alignment(wrap_text=True, vertical="top")
    real, total = data.real_cost_share
    notes.cell(row=len(METHOD_NOTES) + 4, column=1, value=f"{real} de {total} oportunidades usan costo facturado; el resto se estima.")
    notes.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- PDF
def render_pdf(data: ReportData) -> bytes:
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ink, accent, soft, muted = colors.HexColor("#1F2A37"), colors.HexColor("#2F6F5E"), colors.HexColor("#EEF4F1"), colors.HexColor("#5B6770")
    base = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=base["BodyText"], fontName="Helvetica", fontSize=10, leading=14.5, textColor=ink, alignment=TA_LEFT)
    small = ParagraphStyle("small", parent=body, fontSize=8.5, leading=11.5, textColor=muted)
    h1 = ParagraphStyle("h1", parent=body, fontName="Helvetica-Bold", fontSize=22, leading=26, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=13, leading=17, spaceBefore=14, spaceAfter=6)
    cell = ParagraphStyle("cell", parent=body, fontSize=8.5, leading=11)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(muted)
        canvas.drawString(18 * mm, 10 * mm, f"{data.organization} · Confidencial · {data.generated_at:%d-%m-%Y}")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Página {doc.page}")
        canvas.restoreState()

    story: list[Any] = [
        Paragraph("Reporte de desperdicio y ahorro proyectado", h1),
        Paragraph(f"{data.organization} · {long_date(data.generated_at)}", small),
        Spacer(1, 10),
        Paragraph(headline(data), body),
        Spacer(1, 10),
    ]
    kpi = [[usd(data.monthly_spend), usd(data.potential_monthly), usd(data.potential_annual), f"{data.savings_pct:g}%"],
           ["Gasto mensual actual", "Ahorro mensual potencial", "Ahorro anual proyectado", "Del gasto total"]]
    t = Table(kpi, colWidths=[43 * mm] * 4)
    t.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, 0), 12.5),
                           ("TEXTCOLOR", (0, 0), (-1, 0), accent), ("FONTSIZE", (0, 1), (-1, 1), 8), ("TEXTCOLOR", (0, 1), (-1, 1), muted),
                           ("BACKGROUND", (0, 0), (-1, -1), soft), ("TOPPADDING", (0, 0), (-1, 0), 9), ("BOTTOMPADDING", (0, 1), (-1, 1), 8),
                           ("LINEAFTER", (0, 0), (-2, -1), 3, colors.white), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    story.append(t)
    if data.verified:
        story += [Spacer(1, 6), Paragraph(f"Ahorro ya verificado después del despliegue: <b>{usd(data.verified_monthly)}</b> al mes "
                                          f"({len(data.verified)} {'cambio medido' if len(data.verified) == 1 else 'cambios medidos'}).", body)]

    cats = data.grouped("rule")
    if cats:
        story.append(Paragraph("Dónde está el desperdicio", h2))
        d = Drawing(170 * mm, max(40, 18 * len(cats)) * 1.0)
        chart = HorizontalBarChart()
        chart.x, chart.y, chart.width, chart.height = 62 * mm, 6, 100 * mm, max(30, 18 * len(cats) - 8)
        chart.data = [[c[2] for c in reversed(cats)]]
        chart.categoryAxis.categoryNames = [c[0] for c in reversed(cats)]
        chart.categoryAxis.labels.fontName, chart.categoryAxis.labels.fontSize = "Helvetica", 8.5
        chart.valueAxis.labels.fontName, chart.valueAxis.labels.fontSize = "Helvetica", 7.5
        chart.valueAxis.valueMin = 0
        chart.bars[0].fillColor, chart.bars[0].strokeColor = accent, None
        chart.barLabelFormat, chart.barLabels.fontName, chart.barLabels.fontSize = "%0.0f", "Helvetica", 8
        chart.barLabels.nudge = 14
        d.add(chart)
        story.append(d)
        rows = [["Categoría", "Casos", "Ahorro mensual", "Ahorro anual"]] + [[c[0], str(c[1]), usd(c[2]), usd(c[2] * 12)] for c in cats]
        tbl = Table(rows, colWidths=[70 * mm, 20 * mm, 40 * mm, 40 * mm], repeatRows=1)
        tbl.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                                 ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("BACKGROUND", (0, 0), (-1, 0), accent),
                                 ("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, soft]),
                                 ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        story.append(tbl)

    if data.opportunities:
        story.append(Paragraph("Mayores oportunidades", h2))
        top = data.opportunities[:12]
        rows = [["Hallazgo", "Entorno", "Riesgo", "Origen del costo", "Ahorro/mes"]] + [
            [Paragraph(o["title"], cell), o["environment"], o["risk"], Paragraph(source_label(o["cost_source"]), cell), usd(o["monthly_savings"])]
            for o in top]
        tbl = Table(rows, colWidths=[66 * mm, 22 * mm, 17 * mm, 36 * mm, 33 * mm], repeatRows=1)
        tbl.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                                 ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("BACKGROUND", (0, 0), (-1, 0), accent),
                                 ("ALIGN", (-1, 0), (-1, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                 ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, soft]),
                                 ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        story.append(tbl)
        if len(data.opportunities) > len(top):
            story.append(Paragraph(f"Y {len(data.opportunities) - len(top)} más en el archivo Excel.", small))

    envs = data.grouped("environment")
    if envs:
        story.append(Paragraph("Por entorno", h2))
        story.append(Paragraph("; ".join(f"{n}: {usd(s)} al mes ({c} casos)" for n, c, s in envs) + ".", body))

    real, total = data.real_cost_share
    notes = [Paragraph("Cómo leer este reporte", h2)] + [Paragraph("• " + n, body) for n in METHOD_NOTES]
    if total:
        notes.append(Paragraph(f"• En este reporte, {real} de {total} oportunidades usan costo facturado; el resto se estima.", body))
    story.append(KeepTogether(notes))

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=18 * mm,
                      title="Reporte de desperdicio y ahorro proyectado", author="CloudCost Optimizer").build(
        story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
