"""Reporte contra PostgreSQL real: lee las recomendaciones del escaneo demo (con RLS) y genera ambos formatos."""
from __future__ import annotations

import io
import tempfile

from test_lifecycle import ORG, _new_scan, _require_db, _settings  # type: ignore[import-not-found]


def test_report_from_a_real_scan():
    _require_db()
    from cloudcost.db import tenant_tx
    from cloudcost.secrets import SecretResolver
    from cloudcost.services import reports, scan_service
    from openpyxl import load_workbook

    with tenant_tx(ORG) as conn:
        scan_id = _new_scan(conn)
    scan_service.run_scan(ORG, scan_id, settings=_settings(tempfile.mkdtemp()), secrets=SecretResolver())
    with tenant_tx(ORG) as conn:
        data = reports.load_report_data(conn)
    assert data.organization and data.monthly_spend > 0
    assert len(data.opportunities) >= 6
    assert data.opportunities == sorted(data.opportunities, key=lambda o: -o["monthly_savings"])
    assert all(o["cost_source"] for o in data.opportunities)
    assert reports.render_pdf(data).startswith(b"%PDF-")
    wb = load_workbook(io.BytesIO(reports.render_xlsx(data)))
    assert wb["Oportunidades"].max_row == len(data.opportunities) + 1
    assert data.daily_cost, "el escaneo demo registra costos diarios"
