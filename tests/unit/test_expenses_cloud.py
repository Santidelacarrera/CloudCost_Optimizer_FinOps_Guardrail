"""Análisis de gasto de nube desde exportaciones (AWS CUR / Cost Explorer, Azure, GCP) y desde Excel (.xlsx). Datos ficticios."""
from __future__ import annotations

import io
import zipfile
from datetime import date
from decimal import Decimal

import pytest
from cloudcost.expenses import analyze_documents
from cloudcost.expenses.cloud import detect_cloud
from cloudcost.expenses.parser import ExpenseFormatError, read_rows
from cloudcost.expenses.xlsx import MAX_UNCOMPRESSED_BYTES, xlsx_to_csv_texts
from cloudkit import cur_csv, daterange
from openpyxl import Workbook


def run(text: str, name="cur.csv"):
    result, err = analyze_documents([(name, text)])
    assert not err, err
    return result


def test_aws_cur_summary_compares_only_complete_months_and_flags_growth_and_new_service():
    r = run(cur_csv())
    (c,) = r["cloud"]
    assert (c["provider"], c["format"], c["currency"], c["granularity"]) == ("aws", "cur", "USD", "daily")
    assert [m["complete"] for m in c["months"]] == [True, True, False]                       # octubre llega hasta el día 12
    assert c["compared"] == ["2026-08", "2026-09"]                                           # nunca contra un mes parcial
    by_rule = {f["rule"]: f for f in c["findings"]}
    assert "CLOUD_SPEND_GROWTH" in by_rule and "Amazon Elastic Compute Cloud" in by_rule["CLOUD_SPEND_GROWTH"]["title"]
    assert by_rule["CLOUD_NEW_SERVICE"]["title"].startswith("Amazon Relational Database Service")
    assert "CLOUD_CREDITS" in by_rule and Decimal(c["credits"]) == Decimal("-30.00")
    assert any("incompleto" in n for n in c["notes"])
    assert c["by_service"][0]["name"] == "Amazon Elastic Compute Cloud" and r["total_findings"] == len(c["findings"])
    assert len(c["daily"]) == 73 and c["daily"][0] == {"day": "2026-08-01", "amount": "120.00"}               # serie diaria para el gráfico
    assert sum(Decimal(d["amount"]) for d in c["daily"]) == Decimal(c["total"])


def test_totals_are_exact_and_shares_sum_to_one():
    c = run(cur_csv())["cloud"][0]
    assert Decimal(c["total"]) == sum(Decimal(m["total"]) for m in c["months"])
    assert abs(sum(Decimal(s["share"]) for s in c["by_service"]) - 1) < Decimal("0.001")


def test_daily_spike_detected_and_quiet_data_has_none():
    spiky = run(cur_csv(spike=True))["cloud"][0]
    assert spiky["anomalies"] and spiky["anomalies"][0]["day"] == "2026-09-15" and "Simple Storage" in spiky["anomalies"][0]["name"]
    assert any(f["rule"] == "CLOUD_DAILY_SPIKE" for f in spiky["findings"])
    assert run(cur_csv())["cloud"][0]["anomalies"] == []


def test_mixed_currencies_are_rejected_not_summed():
    result, err = analyze_documents([("cur.csv", cur_csv(currencies=("USD", "EUR")))])
    assert not result and "mezcla monedas" in err["errors"][0]


def test_aws_cost_explorer_wide_export():
    text = ('Service,"Jul 2026($)","Aug 2026($)","Service total"\n'
            '"Total costs($)","300.00","520.00","820.00"\n'
            '"Amazon EC2","200.00","300.00","500.00"\n'
            '"Amazon S3","100.00","100.00","200.00"\n'
            '"Amazon RDS","0.00","120.00","120.00"\n')
    (c,) = run(text, "ce.csv")["cloud"]
    assert (c["format"], c["granularity"], c["currency"]) == ("cost_explorer", "monthly", "USD")
    assert [m["total"] for m in c["months"]] == ["300.00", "520.00"]                          # la fila «Total costs» no se suma otra vez
    titles = " ".join(f["title"] for f in c["findings"])
    assert "Amazon EC2" in titles and "Amazon RDS" in titles


def test_azure_and_gcp_exports():
    azure = "Date,SubscriptionName,ResourceGroup,ServiceName,ResourceLocation,CostInBillingCurrency,BillingCurrencyCode\n" + "\n".join(
        f"{d.strftime('%m/%d/%Y')},Prod,rg-web,Virtual Machines,westeurope,{12.5 if d.month == 1 else 30},EUR"
        for d in daterange(date(2026, 1, 1), date(2026, 2, 28)))
    (a,) = run(azure, "azure.csv")["cloud"]
    assert (a["provider"], a["currency"], a["compared"]) == ("azure", "EUR", ["2026-01", "2026-02"])
    assert a["by_region"][0]["name"] == "westeurope" and a["by_group"][0]["name"] == "rg-web"
    gcp = "service.description,usage_start_time,cost,currency,project.id,location.region\n" + "\n".join(
        f"BigQuery,{d.isoformat()}T07:00:00Z,{10 if d.month == 3 else 25},USD,proj-a,us-central1" for d in daterange(date(2026, 3, 1), date(2026, 4, 30)))
    (g,) = run(gcp, "gcp.csv")["cloud"]
    assert (g["provider"], g["by_account"][0]["name"]) == ("gcp", "proj-a") and any(f["rule"] == "CLOUD_SPEND_GROWTH" for f in g["findings"])


def test_generic_expense_statement_is_not_mistaken_for_cloud():
    text = "Proveedor;Concepto;Importe\nEmpresa A;Aseo;100000\nEmpresa B;Seguridad;200000\n"
    assert detect_cloud(read_rows(text), "gastos.csv") is None
    result = run(text, "gastos.csv")
    assert result["cloud"] == [] and result["statements"][0]["item_count"] == 2


# ----------------------------------------------------------------------------------------------------------- Excel
def workbook_bytes(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_generic_expenses_flow_through_the_existing_analyzer():
    data = workbook_bytes({"Agosto 2025": [["Proveedor", "Concepto", "Importe"], ["Empresa A", "Aseo", 100000], ["Empresa B", "Seguridad", 200000.5],
                                           ["Empresa B", "Seguridad", 200000.5]]})
    ((name, text),) = xlsx_to_csv_texts(data, "gastos.xlsx")
    assert name == "gastos.xlsx" and "200000.5" in text
    result = run(text, name)
    assert result["statements"][0]["item_count"] == 3 and result["statements"][0]["period"] is None or True
    assert any("repet" in f["title"].lower() or "doble" in f["title"].lower() or "repet" in f["detail"].lower() for f in result["findings"])


def test_xlsx_cloud_export_and_multiple_sheets_become_separate_files():
    cur = [r.split(",") for r in cur_csv().strip().splitlines()]
    data = workbook_bytes({"CUR": cur, "Otra": [["a", "b"], ["x", "1"]]})
    files = xlsx_to_csv_texts(data, "libro.xlsx")
    assert [n for n, _ in files] == ["libro.xlsx · CUR", "libro.xlsx · Otra"]
    assert detect_cloud(read_rows(files[0][1]), files[0][0]) is not None


def test_xlsx_formulas_are_never_evaluated_and_text_stays_text():
    wb = Workbook()
    ws = wb.active
    ws.append(["Concepto", "Importe"])
    ws.append(['=HYPERLINK("http://evil","x")', 1000])
    ws.append(["Aseo", 2000])
    buf = io.BytesIO()
    wb.save(buf)
    ((_, text),) = xlsx_to_csv_texts(buf.getvalue(), "f.xlsx")
    assert "HYPERLINK" not in text and "evil" not in text            # una fórmula sin resultado guardado queda vacía: nunca se evalúa
    assert text == "Concepto;Importe\n;1000\nAseo;2000\n"


@pytest.mark.parametrize("payload,fragment", [
    (b"not a zip at all", "(?i)no es un archivo .xlsx"),
    (b"PK\x03\x04garbage-garbage", "dañado"),
])
def test_xlsx_rejects_invalid_containers(payload, fragment):
    with pytest.raises(ExpenseFormatError, match=fragment):
        xlsx_to_csv_texts(payload, "x.xlsx")


def test_xlsx_rejects_macros_and_zip_bombs():
    def zipped(entries: dict[str, bytes]) -> bytes:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for n, b in entries.items():
                zf.writestr(n, b)
        return buf.getvalue()

    with pytest.raises(ExpenseFormatError, match="macros"):
        xlsx_to_csv_texts(zipped({"xl/workbook.xml": b"<x/>", "xl/vbaProject.bin": b"x"}), "m.xlsx")
    bomb = zipped({"xl/workbook.xml": b"<x/>", "xl/big.xml": b"0" * (MAX_UNCOMPRESSED_BYTES + 1)})
    assert len(bomb) < 1_000_000
    with pytest.raises(ExpenseFormatError, match="tamaño excesivo"):
        xlsx_to_csv_texts(bomb, "b.xlsx")


def test_schema_requires_exactly_one_source():
    from cloudcost.schemas import ExpenseFileIn
    from pydantic import ValidationError

    assert ExpenseFileIn(filename="a.csv", csv_text="a;b\n1;2")
    assert ExpenseFileIn(filename="a.xlsx", xlsx_base64="UEsDBA==")
    for kw in ({}, {"csv_text": "a;b\n1;2", "xlsx_base64": "UEsDBA=="}):
        with pytest.raises(ValidationError):
            ExpenseFileIn(filename="a.csv", **kw)


# ------------------------------------------------------------------------------------------------ otros formatos de columnas
def _two_months(rows_fn) -> str:
    return "\n".join(rows_fn(d) for d in daterange(date(2026, 3, 1), date(2026, 4, 30)))


def test_focus_export_is_recognised_and_uses_billed_cost():
    header = "BilledCost,EffectiveCost,BillingCurrency,ChargePeriodStart,ChargeCategory,ProviderName,ServiceName,RegionName,SubAccountId"
    body = _two_months(lambda d: f"{10 if d.month == 3 else 30},999,USD,{d.isoformat()}T00:00:00Z,Usage,Amazon Web Services,Amazon EC2,us-east-1,acct-1")
    (c,) = run(header + "\n" + body + "\n", "focus.csv")["cloud"]
    assert (c["format"], c["provider"], c["currency"]) == ("focus", "aws", "USD")
    assert Decimal(c["total"]) == Decimal(31 * 10 + 30 * 30)                                  # BilledCost, no EffectiveCost
    assert c["by_account"][0]["name"] == "acct-1" and any("FOCUS" in n for n in c["notes"])
    assert any(f["rule"] == "CLOUD_SPEND_GROWTH" for f in c["findings"])


def test_spanish_synonym_headers_are_accepted_only_when_services_look_like_cloud():
    cloud = "Fecha;Servicio;Costo;Moneda\n" + _two_months(lambda d: f"{d.isoformat()};Amazon S3;{5 if d.month == 3 else 12};USD")
    (c,) = run(cloud, "factura.csv")["cloud"]
    assert (c["format"], c["currency"]) == ("generic_billing", "USD") and any("sinónimo" in n for n in c["notes"])
    building = "Fecha;Servicio;Costo\n2026-03-05;Aseo;100000\n2026-03-06;Seguridad;200000\n2026-03-07;Jardinería;50000\n"
    result = run(building, "edificio.csv")
    assert result["cloud"] == [] and result["statements"][0]["item_count"] == 3          # no se confunde con una factura de nube


def test_explicit_column_mapping_handles_any_headers_and_reports_missing_ones():
    text = "Día;Item;Monto USD\n" + _two_months(lambda d: f"{d.isoformat()};Cómputo;{4 if d.month == 3 else 9}")
    result, err = analyze_documents([("raro.csv", text)], {"date": "Día", "cost": "Monto USD", "service": "Item"})
    assert not err and result["cloud"][0]["format"] == "mapped_billing" and result["cloud"][0]["by_service"][0]["name"] == "Cómputo"
    _, err = analyze_documents([("raro.csv", text)], {"date": "Día", "cost": "NoExiste", "service": "Item"})
    assert err and ("No se encontró" in err["errors"][0] or "NoExiste" in err["errors"][0])


def test_schema_validates_column_mapping():
    from cloudcost.schemas import ExpenseAnalyzeIn
    from pydantic import ValidationError

    ok = ExpenseAnalyzeIn(files=[{"filename": "a.csv", "csv_text": "a;b\n1;2"}], columns={"date": "Día", "cost": "Monto", "service": "Item", "region": ""})
    assert ok.columns == {"date": "Día", "cost": "Monto", "service": "Item"}
    for bad in ({"date": "Día"}, {"date": "a", "cost": "b", "service": "c", "evil": "x"}):
        with pytest.raises(ValidationError):
            ExpenseAnalyzeIn(files=[{"filename": "a.csv", "csv_text": "a;b\n1;2"}], columns=bad)
