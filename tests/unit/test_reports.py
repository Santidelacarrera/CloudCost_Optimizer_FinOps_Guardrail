"""Informes ejecutivos: formato monetario, Excel con fórmulas vivas, PDF válido y casos vacíos. Sin base de datos."""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

import pytest
from cloudcost.reports.render import cost_source_summary, money, pct, render_pdf, render_xlsx

NOW = datetime(2026, 10, 8, 15, 0, tzinfo=timezone.utc)
LABELS = {"ec2_idle": "Instancias ociosas", "ec2_downsize": "Instancias sobredimensionadas",
          "ebs_orphan": "Volúmenes huérfanos", "snapshot_old": "Snapshots antiguos"}
RISK = {"LOW": "Bajo", "MEDIUM": "Medio", "HIGH": "Alto"}


def _item(n, rule, title, cost, proj, env, risk, src):
    return {"id": f"id{n}", "rule_id": rule, "rule_label": LABELS[rule], "title": title, "status": "PENDING_APPROVAL",
            "status_label": "Pendiente de aprobación", "risk": risk, "risk_label": RISK[risk], "priority": "P1",
            "confidence": 0.9, "destructive": True, "current_monthly_cost": cost, "projected_monthly_cost": proj,
            "estimated_monthly_savings": cost - proj, "created_at": datetime(2026, 10, 1, 12, tzinfo=timezone.utc),
            "resource_id": f"i-{n:04d}", "resource_name": f"svc-{n}", "environment": env, "service": "ec2",
            "region": "us-east-1", "cost_source": src}


def _data(empty: bool = False):
    items = [] if empty else [
        _item(1, "ec2_idle", "Eliminar instancia ociosa <batch> & co", 280.32, 0, "development", "MEDIUM", "cost_explorer"),
        _item(2, "ec2_downsize", "Reducir api-prod: m5.2xlarge → m5.xlarge", 280.32, 140.16, "production", "HIGH", "estimate"),
        _item(3, "ebs_orphan", "Eliminar volumen huérfano vol-0abc", 40.0, 0, "unknown", "MEDIUM", "cost_explorer"),
        _item(4, "snapshot_old", "Eliminar snapshot antiguo snap-1", 22.5, 0, "staging", "LOW", "estimate_upper_bound"),
    ]
    savings = sum(i["estimated_monthly_savings"] for i in items)
    return {
        "organization": "Acme Cloud SpA", "generated_at": NOW, "items": items, "items_truncated": False,
        "kpi": {"monthly_spend": 2400.5, "potential_savings": savings, "savings_pct": 20.1, "recommendations": len(items),
                "pending_approval": len(items), "high_risk": 0 if empty else 1, "verified": 0 if empty else 1, "rejected": 0,
                "realized_savings": 0.0 if empty else 95.0, "expected_savings_verified": 0.0 if empty else 100.0,
                "realization_pct": None if empty else 95.0},
        "by_rule": [] if empty else [{"rule_id": i["rule_id"], "label": i["rule_label"], "count": 1,
                                      "savings": i["estimated_monthly_savings"], "cost": i["current_monthly_cost"]} for i in items],
        "by_environment": [] if empty else [{"environment": i["environment"], "count": 1, "savings": i["estimated_monthly_savings"]} for i in items],
        "by_risk": [] if empty else [{"risk": i["risk"], "label": i["risk_label"], "count": 1, "savings": i["estimated_monthly_savings"]} for i in items],
        "verified": [] if empty else [{"title": "Eliminar instancia vieja", "expected_monthly_savings": 100.0,
                                       "observed_monthly_savings": 95.0, "realization_pct": 95.0, "method": "cost_records",
                                       "window_start": "2026-09-01", "window_end": "2026-09-30", "created_at": NOW}],
    }


def test_money_and_pct_use_spanish_format():
    assert money(1234567.891) == "USD 1.234.567,89"
    assert money(0) == "USD 0,00" and money(None) == "USD 0,00"
    assert pct(28.14) == "28,1 %" and pct(None) == "n/d" and pct(95.0, 0) == "95 %"


def test_cost_source_summary_counts_real_vs_estimated():
    assert cost_source_summary(_data()["items"]) == (2, 2)


def test_xlsx_has_sheets_and_live_formulas():
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(render_xlsx(_data())))                       # sin data_only: vemos las fórmulas
    assert wb.sheetnames == ["Resumen", "Desglose", "Recomendaciones", "Ahorro verificado"]
    res, rec, brk = wb["Resumen"], wb["Recomendaciones"], wb["Desglose"]
    assert res["B7"].value == "=SUM(Recomendaciones!L2:L5)"                 # ahorro potencial: fórmula, no número pegado
    assert res["B8"].value == "=B7*12" and res["B15"].value.startswith("=IF(B14>0")
    assert rec["L2"].value == "=J2-K2" and rec["L6"].value == "=SUM(L2:L5)"
    assert rec.max_row == 6 and rec.auto_filter.ref == "A1:O5" and rec.freeze_panes == "C2"
    assert brk["B3"].value.startswith("=COUNTIF(Recomendaciones!$C$2:$C$5") and "SUMIF" in brk["C3"].value
    assert rec["B2"].value.startswith("Eliminar instancia ociosa <batch> & co")     # el texto se guarda tal cual, sin interpretarse
    assert res["B6"].value == 2400.5 and res["B6"].font.color.rgb.endswith("0000FF")  # dato de entrada en azul


def test_xlsx_formulas_evaluate_without_errors_in_libreoffice(tmp_path):
    import shutil
    import subprocess

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        pytest.skip("LibreOffice no disponible")
    src = tmp_path / "r.xlsx"
    src.write_bytes(render_xlsx(_data()))
    subprocess.run([soffice, "--headless", "--convert-to", "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,1",
                    "--outdir", str(tmp_path), str(src)], check=True, capture_output=True, timeout=120)
    text = next(tmp_path.glob("r*.csv")).read_text(encoding="utf-8")
    assert "Err:" not in text and "#REF" not in text and "#NAME" not in text and "#DIV" not in text
    assert "482.98" in text                                                  # 280.32 + 140.16 + 40 + 22.5


def test_xlsx_empty_report_does_not_break():
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(render_xlsx(_data(empty=True))))
    assert wb["Resumen"]["B10"].value == 0 and wb["Recomendaciones"].max_row == 1


def test_pdf_is_valid_and_paginates():
    pdf = render_pdf(_data())
    assert pdf.startswith(b"%PDF-") and pdf.rstrip().endswith(b"%%EOF")
    assert pdf.count(b"/Type /Page\n") + pdf.count(b"/Type /Page ") >= 1
    many = _data()
    many["items"] = [_item(i, "ebs_orphan", f"Eliminar volumen huérfano vol-{i:05d}", 40.0, 0, "production", "MEDIUM", "estimate")
                     for i in range(300)]
    assert len(render_pdf(many)) > len(pdf)                                  # el detalle largo no rompe el generador


def test_pdf_empty_and_special_characters():
    assert render_pdf(_data(empty=True)).startswith(b"%PDF-")
    data = _data()
    data["organization"] = "Ñandú <b>&</b> Cía"
    assert render_pdf(data).startswith(b"%PDF-")                             # el marcado en datos no rompe Paragraph
